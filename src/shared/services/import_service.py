"""Bulk import service for categories and products from Excel files."""

from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from database.db import DatabaseError
from database.db_manager import get_db
from shared.number_utils import to_decimal
from shared.record_utils import get_record_value
from utils.database import (
    create_category_db,
    create_inventory_row_db,
    create_product_db,
    get_category_by_name_db,
    get_product_by_name_db,
    insert_inventory_movement,
    update_category_db,
    update_product_db,
    _get_inventory_row,
)
from utils.logger_factory import get_api_logger

logger = get_api_logger(__name__)

_REQUIRED_CATEGORY_FIELDS = {"codigo", "nombre", "descripcion"}
_REQUIRED_PRODUCT_FIELDS = {"nombre", "descripcion", "precio", "category_name"}


def _normalize_string(value: Any) -> str:
    """Return a stripped string or empty string for None."""
    if value is None:
        return ""
    return str(value).strip()


def _log_processing_error(
    upload_id: int,
    row_number: int,
    column_name: str | None,
    error_type: str,
    error_message: str,
    raw_data: str | None = None,
) -> None:
    """Insert a processing error record for a failed row."""
    try:
        db = get_db()
        db.execute(
            """
            INSERT INTO csv_processing_errors
            (upload_id, row_number, column_name, error_type, error_message, raw_data)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (upload_id, row_number, column_name, error_type, error_message, raw_data),
        )
    except DatabaseError as e:
        logger.error("Error registrando fallo de procesamiento: %s", e)


def _parse_positive_int(value: Any, field_name: str) -> int:
    """Parse a non-negative integer or raise ValueError."""
    try:
        parsed = int(float(_normalize_string(value)))
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{field_name} debe ser un número entero") from exc
    if parsed < 0:
        raise ValueError(f"{field_name} no puede ser negativo")
    return parsed


def _parse_positive_decimal(value: Any, field_name: str) -> Decimal:
    """Parse a non-negative decimal or raise ValueError."""
    try:
        parsed = to_decimal(_normalize_string(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field_name} debe ser un número válido") from exc
    if parsed < 0:
        raise ValueError(f"{field_name} no puede ser negativo")
    return parsed


def _load_workbook_sheet(file_path: str, sheet_name: str) -> Any:
    """Load a workbook and return the requested sheet.

    The caller is responsible for closing the returned workbook.
    """
    workbook = load_workbook(filename=file_path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        workbook.close()
        raise ValueError(f"No se encontró la hoja '{sheet_name}'. Hojas disponibles: {available}")
    return workbook, workbook[sheet_name]


def _row_to_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    """Build a dict from an openpyxl row and headers."""
    return {
        str(header).strip().lower(): value
        for header, value in zip(headers, row)
        if header is not None
    }


def import_categories_from_xlsx(file_path: str, upload_id: int) -> dict[str, Any]:
    """Import or update categories from the 'categories' sheet of an Excel file.

    Expected columns: codigo, nombre, descripcion, parent_name (optional).

    Returns a summary dict with processed, created, updated and error counts.
    """
    workbook, sheet = _load_workbook_sheet(file_path, "categories")
    try:
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return {"processed_rows": 0, "created": 0, "updated": 0, "errors": 0}

    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    data_rows = rows[1:]

    missing_required = _REQUIRED_CATEGORY_FIELDS - set(headers)
    if missing_required:
        raise ValueError(f"Columnas requeridas faltantes: {', '.join(sorted(missing_required))}")

    created = 0
    updated = 0
    errors = 0
    processed = 0

    for row_index, row in enumerate(data_rows, start=2):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        processed += 1
        data = _row_to_dict(headers, row)

        try:
            codigo = _normalize_string(data.get("codigo")).upper()
            nombre = _normalize_string(data.get("nombre"))
            descripcion = _normalize_string(data.get("descripcion"))
            parent_name = _normalize_string(data.get("parent_name")) or None

            if not codigo:
                raise ValueError("codigo no puede estar vacío")
            if not nombre:
                raise ValueError("nombre no puede estar vacío")
            if not descripcion:
                raise ValueError("descripcion no puede estar vacío")

            parent_id = None
            if parent_name:
                parent = get_category_by_name_db(parent_name)
                if not parent:
                    raise ValueError(f"categoría padre no encontrada: {parent_name}")
                parent_id = int(get_record_value(parent, "id", fallback_index=0))

            existing = get_category_by_name_db(nombre)
            if existing:
                category_id = int(get_record_value(existing, "id", fallback_index=0))
                update_data = {
                    "codigo": codigo,
                    "nombre": nombre,
                    "descripcion": descripcion,
                    "parent_id": parent_id,
                }
                if update_category_db(category_id, update_data):
                    updated += 1
                else:
                    raise ValueError("no se pudo actualizar la categoría")
            else:
                new_id = create_category_db(codigo, nombre, descripcion, parent_id)
                if new_id is None:
                    raise ValueError("no se pudo crear la categoría")
                created += 1

        except ValueError as exc:
            errors += 1
            _log_processing_error(
                upload_id=upload_id,
                row_number=row_index,
                column_name=None,
                error_type="VALIDATION_ERROR",
                error_message=str(exc),
                raw_data=str(row),
            )
        except Exception as exc:
            errors += 1
            logger.exception("Error inesperado procesando categoría fila %s", row_index)
            _log_processing_error(
                upload_id=upload_id,
                row_number=row_index,
                column_name=None,
                error_type="UNEXPECTED_ERROR",
                error_message=str(exc),
                raw_data=str(row),
            )

    return {
        "processed_rows": processed,
        "created": created,
        "updated": updated,
        "errors": errors,
    }


def import_products_from_xlsx(file_path: str, upload_id: int) -> dict[str, Any]:
    """Import or update products from the 'products' sheet of an Excel file.

    Expected columns: nombre, descripcion, precio, category_name, stock_inicial (optional).

    For existing products, the stock is adjusted via an inventory movement of type
    'ajuste' so that stock_actual matches the stock_inicial value from the file.

    Returns a summary dict with processed, created, updated and error counts.
    """
    workbook, sheet = _load_workbook_sheet(file_path, "products")
    try:
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return {"processed_rows": 0, "created": 0, "updated": 0, "errors": 0}

    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    data_rows = rows[1:]

    missing_required = _REQUIRED_PRODUCT_FIELDS - set(headers)
    if missing_required:
        raise ValueError(f"Columnas requeridas faltantes: {', '.join(sorted(missing_required))}")

    created = 0
    updated = 0
    errors = 0
    processed = 0

    for row_index, row in enumerate(data_rows, start=2):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        processed += 1
        data = _row_to_dict(headers, row)

        try:
            nombre = _normalize_string(data.get("nombre"))
            descripcion = _normalize_string(data.get("descripcion"))
            category_name = _normalize_string(data.get("category_name"))

            if not nombre:
                raise ValueError("nombre no puede estar vacío")
            if not descripcion:
                raise ValueError("descripcion no puede estar vacío")
            if not category_name:
                raise ValueError("category_name no puede estar vacío")

            precio = _parse_positive_decimal(data.get("precio"), "precio")
            stock_inicial = 0
            if "stock_inicial" in headers and data.get("stock_inicial") not in (None, ""):
                stock_inicial = _parse_positive_int(data.get("stock_inicial"), "stock_inicial")

            category = get_category_by_name_db(category_name)
            if not category:
                raise ValueError(f"categoría no encontrada: {category_name}")
            category_id = int(get_record_value(category, "id", fallback_index=0))

            existing = get_product_by_name_db(nombre)
            if existing:
                product_id = int(get_record_value(existing, "id", fallback_index=0))
                update_data = {
                    "nombre": nombre,
                    "descripcion": descripcion,
                    "precio": float(precio),
                    "category_id": category_id,
                }
                if not update_product_db(product_id, update_data):
                    raise ValueError("no se pudo actualizar el producto")

                _ensure_stock_matches(product_id, stock_inicial, upload_id, row_index)
                updated += 1
            else:
                new_id = create_product_db(nombre, descripcion, float(precio), category_id)
                if new_id is None:
                    raise ValueError("no se pudo crear el producto")
                if not create_inventory_row_db(new_id, stock_inicial):
                    raise ValueError("no se pudo crear la fila de inventario")
                created += 1

        except ValueError as exc:
            errors += 1
            _log_processing_error(
                upload_id=upload_id,
                row_number=row_index,
                column_name=None,
                error_type="VALIDATION_ERROR",
                error_message=str(exc),
                raw_data=str(row),
            )
        except Exception as exc:
            errors += 1
            logger.exception("Error inesperado procesando producto fila %s", row_index)
            _log_processing_error(
                upload_id=upload_id,
                row_number=row_index,
                column_name=None,
                error_type="UNEXPECTED_ERROR",
                error_message=str(exc),
                raw_data=str(row),
            )

    return {
        "processed_rows": processed,
        "created": created,
        "updated": updated,
        "errors": errors,
    }


def _ensure_stock_matches(
    product_id: int,
    target_stock: int,
    upload_id: int,
    row_index: int,
) -> None:
    """Adjust inventory so stock_actual equals target_stock for an existing product.

    If the product has no inventory row, one is created with the target stock.
    Otherwise an 'ajuste' movement is recorded with the difference.
    """
    inventory_row = _get_inventory_row(product_id)
    if inventory_row is None:
        if not create_inventory_row_db(product_id, target_stock):
            raise ValueError("no se pudo crear la fila de inventario")
        return

    current_stock = int(get_record_value(inventory_row, "stock_actual", fallback_index=0) or 0)
    difference = target_stock - current_stock
    if difference == 0:
        return

    try:
        movement_id = insert_inventory_movement(
            product_id=product_id,
            tipo_movimiento="ajuste",
            cantidad=difference,
            referencia=f"import_upload_{upload_id}",
        )
    except DatabaseError as exc:
        raise ValueError(f"no se pudo ajustar el stock: {exc}") from exc

    if movement_id is None:
        raise ValueError("no se pudo registrar el movimiento de ajuste de stock")
